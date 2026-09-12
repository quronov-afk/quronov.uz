"""Foto galereya ro'yxatini yasaydi: izoh, tartib va guruhlash.

Izohlar «Foto/quronov-foto-izohlar.docx» jadvalidan (f001–f099) hamda
«Foto/01» papkasidagi fayl nomlaridan (f100 dan boshlab) olinadi.

Tartib qoidasi:
  1) Dilmurod Quronovning ikki surati bilan boshlanadi;
  2) rasmiy suratlar — uchrashuv, taqdimot, konferensiya, ijodiy jamoalar;
  3) safar, sayohat va do'stona suratlar;
  4) oila va maishiy suratlar eng oxirida.
Har guruh ichida suratlar shakliga (yotiq/tik) qarab saralanadi va bir
sujetdagi suratlar yonma-yon tushmaydi.

Ishlatish:  python3 tools/rasmlardan.py
Natija:     data/rasmlar.json
"""

import json, re, sys
from pathlib import Path

from PIL import Image
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parent))
from kitobdan import apostrof

ROOT = Path(__file__).resolve().parent.parent
FOTO = ROOT / "site" / "img" / "foto"
IZOH_FAYL = ROOT / "Foto" / "quronov-foto-izohlar.docx"

# «Foto/01» papkasidagi suratlar: izoh fayl nomining o'zida yozilgan
YANGILAR = {
    "f100": "«Jajji akademik» jurnali jamoasi", "f101": "«Jajji akademik» jurnali jamoasi",
    "f102": "«Jajji akademik» jurnalida", "f103": "«Sharq yulduzi» jurnali jamoasi",
    "f104": "«Sharq yulduzi» jurnali jamoasi", "f105": "«Yoshlik» jurnali jamoasi",
    "f106": "Abdulla Oripov bilan",
    "f107": "Akademik Naim Karimov va professor Ulugʻbek Hamdam bilan",
    "f108": "Bolalik", "f109": "Erkin Vohidov, Ulugʻbek Hamdam, Isajon Sulton bilan",
    "f110": "Ijodiy uchrashuv", "f111": "Ijodiy uchrashuv", "f112": "Ijodkor yoshlar",
    "f113": "Bolalar adabiyoti dargʻalari: Anvar Obidjon, Xudoyberdi Toʻxtaboyev, Tursunboy Adashboyev, bastakor Shermat Yormatov, noshir Sanjar Nazar bilan", "f114": "Imzo marosimi", "f115": "Imzo marosimi",
    "f116": "Kitob taqdimoti", "f117": "Kitob taqdimotida", "f118": "Kursdoshlar",
    "f119": "Professor Qozoqboy Yoʻldosh bilan", "f120": "RTM jamoasi",
    "f121": "Shuhrat Sattorov bilan", "f122": "Umarali Normatov xonadonida",
    "f123": "Xurshid Davron bilan",
}

# izohlardagi terish xatolari
TUZATISH = {
    "Dilmurod Quronob": "Dilmurod Quronov",
    "adminstratsiyasi": "administratsiyasi",
    "Mehrinov Abbosova": "Mehrinoz Abbosova",
    "loyihasing": "loyihasining",
    "mualliflari jamosi": "mualliflari jamoasi",
    "yozvchilari": "yozuvchilari",
    "Yangi Toshentdagi": "Yangi Toshkentdagi",
    "Dyublaj": "Dublyaj",
    "Olmoniya": "Germaniya",
    "Muhammadjon Xoʻjayevlar": "Muhammadjon Xoʻjayev",
    "asari tarjimoni (ingliz tiliga) Nick Walmsley":
        "asarining ingliz tiliga tarjimoni Nick Walmsley bilan",
}

# suratlar guruhi: rasmiy (2), safar va doʻstona (3), oila va maishiy (4)
OILA = ["oila", "onam bilan", "ona bilan", "nabira", "qarindosh", "bolalik", "jigar",
        "xonadonida", "kursdosh", "tabiat qoʻynida"]

# rasmiy suratlar ichida eng oldinga chiqadiganlar
YUQORI = ["prezident", "administratsiya", "senator", "vazir", "mirziyoyeva"]
SAFAR = ["germaniya", "italiya", "berlin", "rim", "bolonya", "florensiya", "muzey",
         "playstation", "futbol", "nihol ekish", "doʻstlar", "qadrdon", "davrasida",
         "tabiat", "muxlis"]

# galereyaga chiqmaydigan suratlar
OCHIRILGAN = {"f002"}

BOSH = ["f042", "f057"]          # galereya shu ikki suratdan boshlanadi
DQ_SOCHMA = ["f046", "f001", "f061", "f045"]   # D. Quronovning qolgan suratlari


def izohlar():
    """f001–f099 uchun jadvaldan, f100 dan keyingilar uchun fayl nomidan."""
    natija = dict(YANGILAR)
    t = Document(str(IZOH_FAYL)).tables[0]
    for qator in t.rows[1:]:
        belgi = re.search(r"f\d{3}", qator.cells[0].text)
        if not belgi:
            continue
        bolaklar = [p.text.strip() for p in qator.cells[2].paragraphs if p.text.strip()]
        matn = bolaklar[0] if bolaklar else ""
        for keyingi in bolaklar[1:]:           # jadvalda yangi satrdan yozilgan ismlar
            ayirgich = " " if matn.endswith((":", ",", "–", "-")) else ": "
            matn += ayirgich + keyingi
        natija[belgi.group()] = matn
    return {k: tozalash(v) for k, v in natija.items()}


def tozalash(s):
    s = re.sub(r"\s+", " ", s).strip(" .")
    for xato, togri in TUZATISH.items():
        s = s.replace(xato, togri)
    s = apostrof(s)
    return s.replace("“", "«").replace("”", "»")


def guruh(nom, izoh):
    """Surat qaysi bo'limga tegishli: 1 – D. Quronov, 2 – rasmiy, 3 – safar, 4 – oila."""
    if nom in BOSH or nom in DQ_SOCHMA:
        return 1
    past = izoh.lower()
    if any(k in past for k in OILA):
        return 4
    if any(k in past for k in SAFAR):
        return 3
    return 2


def shakl(nom):
    w, h = Image.open(FOTO / f"{nom}.jpg").size
    return 0 if w > h * 1.15 else (1 if abs(w - h) <= h * 0.15 else 2)


def yoyish(royxat, sujet):
    """Bir sujetdagi (bir xil izohli) suratlar ketma-ket tushmasin."""
    qolgan, natija = list(royxat), []
    while qolgan:
        oxirgi = sujet.get(natija[-1]) if natija else None
        keyingi = next((n for n in qolgan if sujet.get(n) != oxirgi), qolgan[0])
        natija.append(keyingi)
        qolgan.remove(keyingi)
    return natija


def main():
    izoh = izohlar()
    suratlar = sorted(p.stem for p in FOTO.glob("f*.jpg") if p.stem not in OCHIRILGAN)
    yoq = [n for n in suratlar if n not in izoh]
    if yoq:
        print("izohsiz:", ", ".join(yoq))

    # bir xil izoh — bir sujet; izohsizlari alohida sanaladi
    sujet = {n: (izoh.get(n) or n).lower() for n in suratlar}

    tartib = list(BOSH)
    for g in (2, 3, 4):
        uniki = [n for n in suratlar
                 if n not in BOSH and n not in DQ_SOCHMA and guruh(n, izoh.get(n, "")) == g]

        def kalit(n):
            past = izoh.get(n, "").lower()
            return (0 if any(k in past for k in YUQORI) else 1, shakl(n))

        tartib += yoyish(sorted(uniki, key=kalit), sujet)
    for k, nom in enumerate(DQ_SOCHMA):        # D. Quronovning suratlari oraga sochiladi
        tartib.insert(7 + k * 11, nom)
    tartib += [n for n in suratlar if n not in tartib]

    yopishgan = [(a, b) for a, b in zip(tartib, tartib[1:]) if sujet[a] == sujet[b]]
    print(f"{len(tartib)} ta surat | yonma-yon oʻxshash: {len(yopishgan)}")

    (ROOT / "data" / "rasmlar.json").write_text(
        json.dumps([{"fayl": f"img/foto/{n}.jpg", "nom": izoh.get(n, "")}
                    for n in tartib], ensure_ascii=False, indent=1) + "\n",
        encoding="utf-8")


if __name__ == "__main__":
    main()
